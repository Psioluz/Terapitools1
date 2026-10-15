import sys, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment

out = sys.argv[1]
test = len(sys.argv) > 2 and sys.argv[2] == "test"

F = "Arial"
def font(**k): return Font(name=F, **k)
HDR = PatternFill("solid", fgColor="1F3864")
SUB = PatternFill("solid", fgColor="D9E1F2")
INP = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BLUE = font(color="0000FF")
GREEN = font(color="008000")

N0, N1 = 5, 734          # filas de datos
SKUS = [(153523, "CAJA MEDIANA 22X33X30 411 C"),
        (153553, "CAJA 1/24 PEQUEÑA 11X22X15 312 C IMP"),
        (153526, "CAJA 1/2 GRANDE 33X44X30 411 C IMP")]

wb = Workbook()
def title(ws, text, sub):
    ws["A1"] = text; ws["A1"].font = font(bold=True, size=14, color="1F3864")
    ws["A2"] = sub; ws["A2"].font = font(italic=True, size=9, color="595959")
def header(ws, row, labels, col0=1):
    for i, l in enumerate(labels):
        c = ws.cell(row=row, column=col0 + i, value=l)
        c.font = font(bold=True, color="FFFFFF"); c.fill = HDR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BOX

# ---------------- Instrucciones ----------------
ws = wb.active; ws.title = "Instrucciones"
title(ws, "Control de stock de cajas (SKU 153523 / 153553 / 153526)", "Fuente: correo 'PE4 RSP - TOTAL INVENTORY SNAPSHOT' (yobelrsp@yobel.biz), hoja BaseDeDatosReporteador, Atributo 2 = DISPONIBLE")
lines = [
 ("Hojas", True),
 ("Config: parámetros del proceso (celdas amarillas = editables): SKUs, filtro, remitente/asunto, horarios, destinatarios de alerta.", False),
 ("Stock_Diario: una fila por día con el stock DISPONIBLE de cada SKU. La llena el script automáticamente (no editar a mano salvo correcciones).", False),
 ("Recepciones: (opcional) registre aquí las recepciones conocidas. Si un día sube el stock y no hay recepción registrada, se asume recepción = aumento.", False),
 ("Consumo_Diario: consumo por intervalo = Stock anterior + Recepción - Stock actual (fórmulas, no editar).", False),
 ("Indicador: stock actual, consumo promedio diario (últimos N días y mes en curso), días de cobertura, y tabla de consumo/promedio mensual.", False),
 ("Log: bitácora de ejecuciones del script (cargas, correos no encontrados, alertas enviadas).", False),
 ("", False),
 ("Cálculo", True),
 ("Stock del día = suma de 'Saldo' de las filas con Código Item = SKU y Atributo 2 = DISPONIBLE (un SKU puede tener varias filas/ubicaciones).", False),
 ("Consumo del intervalo = Stock día anterior + Recepciones del intervalo - Stock del día (mínimo 0).", False),
 ("Promedio diario = consumo del periodo / días transcurridos del periodo (corrige fines de semana o días sin archivo).", False),
 ("Días de cobertura = Stock actual / consumo promedio diario de los últimos N días.", False),
 ("", False),
 ("Instalación: ver README.md (Google Sheets + Apps Script, revisiones automáticas y alerta por correo si no llega el archivo).", True),
]
for i, (t, b) in enumerate(lines, start=4):
    ws.cell(row=i, column=1, value=t).font = font(bold=b)
ws.column_dimensions["A"].width = 140

# ---------------- Config ----------------
cf = wb.create_sheet("Config")
title(cf, "Configuración", "Celdas amarillas editables. El script lee estos valores por nombre (rangos con nombre cfg_*).")
header(cf, 3, ["Parámetro", "Valor", "Nota"])
params = [
 ("cfg_Remitente", "Remitente del correo", "yobelrsp@yobel.biz", "Dirección del bot (ver captura del correo)."),
 ("cfg_Asunto", "Asunto contiene", "TOTAL INVENTORY SNAPSHOT", "Texto fijo del asunto 'PE4 RSP - TOTAL INVENTORY SNAPSHOT (dd/mm/aaaa hh:mm:ss)'."),
 ("cfg_PatronAdjunto", "Nombre adjunto empieza con", "TotalFotoRsp", "Ej.: TotalFotoRsp_-_2026101508.xlsx"),
 ("cfg_HojaOrigen", "Hoja del archivo", "BaseDeDatosReporteador", ""),
 ("cfg_ColCodigo", "Columna código", "Código Item", "Encabezado exacto en el archivo."),
 ("cfg_ColFiltro", "Columna filtro", "Atributo 2", ""),
 ("cfg_ValorFiltro", "Valor filtro", "DISPONIBLE", "Solo se suman filas con este valor."),
 ("cfg_ColCantidad", "Columna cantidad", "Saldo", ""),
 ("cfg_HorasRevision", "Horas de revisión", "08:45, 09:15, 09:45, 10:15, 12:00", "El correo llega ~08:30. El script revisa a estas horas (±15 min, límite de Google)."),
 ("cfg_HoraLimite", "Hora límite", "10:00", "Si a esta hora no se cargó el stock del día, se envía la alerta (una vez por día)."),
 ("cfg_Destinatarios", "Destinatarios alerta", "correo1@rashperu.com, correo2@rashperu.com", "Separar con coma. REEMPLAZAR por los correos reales."),
 ("cfg_FinesSemana", "Alertar sábados y domingos", "SI", "SI / NO. Poner NO si el bot no envía en fines de semana."),
 ("cfg_DiasPromedio", "Días para promedio móvil", 7, "Ventana del consumo promedio diario usado para la cobertura."),
 ("cfg_CoberturaMin", "Cobertura mínima (días)", 15, "Debajo de este valor el Indicador marca REPONER y el script envía aviso."),
 ("cfg_AvisoCobertura", "Enviar aviso de cobertura baja", "SI", "SI / NO"),
 ("cfg_CarpetaDrive", "ID carpeta Drive para archivar adjuntos", "", "Opcional. Vacío = no se archiva el .xlsx original."),
]
r = 4
for name, lab, val, note in params:
    cf.cell(row=r, column=1, value=lab).font = font()
    c = cf.cell(row=r, column=2, value=val); c.font = BLUE; c.fill = INP; c.border = BOX
    cf.cell(row=r, column=3, value=note).font = font(size=9, color="595959")
    wb.defined_names[name] = DefinedName(name, attr_text=f"Config!$B${r}")
    r += 1
dv = DataValidation(type="list", formula1='"SI,NO"', allow_blank=False)
cf.add_data_validation(dv)
for rr in range(4, r):
    if cf.cell(row=rr, column=2).value in ("SI", "NO"): dv.add(cf.cell(row=rr, column=2))
sku_row = r + 2
cf.cell(row=sku_row - 1, column=1, value="SKUs a extraer (orden fijo: columnas E, F, G de Stock_Diario)").font = font(bold=True)
header(cf, sku_row, ["SKU", "Descripción", ""])
for i, (s, d) in enumerate(SKUS):
    rr = sku_row + 1 + i
    c = cf.cell(row=rr, column=1, value=s); c.font = BLUE; c.fill = INP; c.border = BOX
    c = cf.cell(row=rr, column=2, value=d); c.font = BLUE; c.fill = INP; c.border = BOX
SKU_R = [sku_row + 1 + i for i in range(3)]
wb.defined_names["cfg_SKUs"] = DefinedName("cfg_SKUs", attr_text=f"Config!$A${SKU_R[0]}:$A${SKU_R[2]}")
cf.column_dimensions["A"].width = 38; cf.column_dimensions["B"].width = 45; cf.column_dimensions["C"].width = 80
CFG_DIAS = "Config!$B$16"; CFG_COB = "Config!$B$17"
assert cf["A16"].value == "Días para promedio móvil" and cf["A17"].value == "Cobertura mínima (días)"

# ---------------- Stock_Diario ----------------
sd = wb.create_sheet("Stock_Diario")
title(sd, "Stock diario DISPONIBLE (lo llena el script)", "Una fila por fecha de snapshot, ordenada ascendente. No insertar ni borrar filas: solo valores.")
header(sd, 4, ["Fecha", "Hora snapshot", "Archivo", "Origen", None, None, None, "Filas encontradas", "Cargado el"])
for i, col in enumerate("EFG"):
    sd[f"{col}4"] = f'="Stock "&Config!$A${SKU_R[i]}'
    sd[f"{col}3"] = f"=Config!$B${SKU_R[i]}"; sd[f"{col}3"].font = font(size=8, color="008000")
    sd[f"{col}3"].alignment = Alignment(wrap_text=True)
seed = [(dt.date(2026, 10, 15), "08:29", "TotalFotoRsp_-_2026101508.xlsx", "Carga inicial", 6725, 5025, 11025, "1/2/1", dt.datetime(2026, 10, 15, 8, 45))]
if test:
    seed += [
      (dt.date(2026,10,16),"08:30","t","test",6500,4900,10800,"",None),
      (dt.date(2026,10,17),"08:30","t","test",6300,4800,12500,"",None),   # 153526 recepción no registrada
      (dt.date(2026,10,20),"08:30","t","test",7000,4500,12000,"",None),   # 153523 recepción registrada 1000 (cons 300), 3 días
      (dt.date(2026,11,2),"08:30","t","test",6800,4400,11900,"",None),
    ]
for i, row in enumerate(seed):
    for j, v in enumerate(row):
        sd.cell(row=N0 + i, column=1 + j, value=v)
for rr in range(N0, N1 + 1):
    sd.cell(row=rr, column=1).number_format = "dd/mm/yyyy"
    sd.cell(row=rr, column=9).number_format = "dd/mm/yyyy hh:mm"
    for col in (5, 6, 7): sd.cell(row=rr, column=col).number_format = "#,##0"
    for col in range(1, 10): sd.cell(row=rr, column=col).font = font()
sd.freeze_panes = "A5"
for col, w in zip("ABCDEFGHI", [12, 13, 34, 14, 16, 16, 16, 16, 17]): sd.column_dimensions[col].width = w
sd.row_dimensions[3].height = 24

# ---------------- Recepciones ----------------
rc = wb.create_sheet("Recepciones")
title(rc, "Recepciones (opcional, ingreso manual)", "Ejemplo de formato: 20/10/2026 | 153523 | 1000 | GR-0001-123 | Ingreso proveedor. Fecha = día en que la recepción ya figura en el stock del bot.")
header(rc, 4, ["Fecha", "SKU", "Cantidad", "Guía / Documento", "Comentario"])
dvs = DataValidation(type="list", formula1=f"=Config!$A${SKU_R[0]}:$A${SKU_R[2]}", allow_blank=True)
rc.add_data_validation(dvs); dvs.add(f"B5:B2000")
if test:
    rc["A5"] = dt.date(2026,10,20); rc["B5"] = 153523; rc["C5"] = 1000
for rr in range(5, 2001):
    rc.cell(row=rr, column=1).number_format = "dd/mm/yyyy"
    for col in range(1, 6):
        c = rc.cell(row=rr, column=col); c.font = BLUE
        if rr <= 60: c.fill = INP
rc.freeze_panes = "A5"
for col, w in zip("ABCDE", [12, 12, 12, 22, 40]): rc.column_dimensions[col].width = w

# ---------------- Consumo_Diario ----------------
cd = wb.create_sheet("Consumo_Diario")
title(cd, "Consumo por intervalo entre snapshots", "Consumo = Stock anterior + Recepción - Stock actual. Si el stock sube sin recepción registrada, se asume Recepción = aumento.")
cd["A4"] = "Fecha"; cd["B4"] = "Días"
labels = ["Stock", "Recepción", "Consumo"]
blocks = [("C", "D", "E"), ("F", "G", "H"), ("I", "J", "K")]
SDCOL = ["E", "F", "G"]
for i, (s, rcol, c) in enumerate(blocks):
    cd.merge_cells(f"{s}3:{c}3")
    cd[f"{s}3"] = f'=Config!$A${SKU_R[i]}&" - "&Config!$B${SKU_R[i]}'
    cd[f"{s}3"].font = font(bold=True, color="FFFFFF"); cd[f"{s}3"].fill = PatternFill("solid", fgColor="2F5597")
    cd[f"{s}3"].alignment = Alignment(horizontal="center")
    for k, col in enumerate((s, rcol, c)): cd[f"{col}4"] = labels[k]
cd["L4"] = "Consumo total"; cd["M4"] = "Observación"
for col in "ABCDEFGHIJKLM":
    c = cd[f"{col}4"]; c.font = font(bold=True, color="FFFFFF"); c.fill = HDR; c.border = BOX
    c.alignment = Alignment(horizontal="center", wrap_text=True)
RECSUM = 'SUMIFS(Recepciones!$C$5:$C$2000,Recepciones!$B$5:$B$2000,Config!$A${sr},Recepciones!$A$5:$A$2000,">"&$A{p},Recepciones!$A$5:$A$2000,"<="&$A{r})'
for r in range(N0, N1 + 1):
    p = r - 1
    cd[f"A{r}"] = f'=IF(Stock_Diario!A{r}="","",Stock_Diario!A{r})'
    cd[f"B{r}"] = '=""' if r == N0 else f'=IF($A{r}="","",$A{r}-$A{p})'
    for i, (s, rcol, c) in enumerate(blocks):
        cd[f"{s}{r}"] = f'=IF($A{r}="","",Stock_Diario!{SDCOL[i]}{r})'
        if r == N0:
            cd[f"{rcol}{r}"] = '=""'; cd[f"{c}{r}"] = '=""'
        else:
            rs = RECSUM.format(sr=SKU_R[i], p=p, r=r)
            cd[f"{rcol}{r}"] = f'=IF($A{r}="","",IF({rs}>0,{rs},MAX(0,{s}{r}-{s}{p})))'
            cd[f"{c}{r}"] = f'=IF($A{r}="","",MAX(0,{s}{p}+{rcol}{r}-{s}{r}))'
    if r == N0:
        cd[f"L{r}"] = '=""'; cd[f"M{r}"] = f'=IF($A{r}="","","Primer registro (base)")'
    else:
        cd[f"L{r}"] = f'=IF($A{r}="","",E{r}+H{r}+K{r})'
        anyrec = f'SUMIFS(Recepciones!$C$5:$C$2000,Recepciones!$A$5:$A$2000,">"&$A{p},Recepciones!$A$5:$A$2000,"<="&$A{r})'
        cd[f"M{r}"] = (f'=IF($A{r}="","",IF(AND(D{r}+G{r}+J{r}>0,{anyrec}=0),"Aumento de stock sin recepción registrada",'
                       f'IF(B{r}>1,"Intervalo de "&B{r}&" días","")))')
    cd[f"A{r}"].number_format = "dd/mm/yyyy"
    for col in "CDEFGHIJKL": cd[f"{col}{r}"].number_format = "#,##0"
    for col in "ABCDEFGHIJKLM": cd[f"{col}{r}"].font = font()
cd.freeze_panes = "C5"
for col, w in zip("ABCDEFGHIJKLM", [12, 6] + [11] * 9 + [13, 42]): cd.column_dimensions[col].width = w
cd.conditional_formatting.add(f"D{N0}:D{N1}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="E2EFDA")))
cd.conditional_formatting.add(f"G{N0}:G{N1}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="E2EFDA")))
cd.conditional_formatting.add(f"J{N0}:J{N1}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="E2EFDA")))

# ---------------- Indicador ----------------
ind = wb.create_sheet("Indicador")
title(ind, "Indicador de consumo de cajas", "Se actualiza solo con cada carga. Cobertura = stock actual / consumo promedio diario de los últimos N días (Config).")
FD = f"Consumo_Diario!$A${N0}:$A${N1}"; DD = f"Consumo_Diario!$B${N0}:$B${N1}"
ind["A3"] = "Último snapshot:"; ind["B3"] = f'=IF(COUNT(Stock_Diario!$A${N0}:$A${N1})=0,"Sin datos",MAX(Stock_Diario!$A${N0}:$A${N1}))'
ind["B3"].number_format = "dd/mm/yyyy"
ind["C3"] = "Carga de hoy:"; ind["D3"] = f'=IF(COUNTIF(Stock_Diario!$A${N0}:$A${N1},TODAY())>0,"OK - cargado","PENDIENTE")'
ind["E3"] = "N registros:"; ind["F3"] = f"=COUNT(Stock_Diario!$A${N0}:$A${N1})"
for c in ("A3", "C3", "E3"): ind[c].font = font(bold=True)
for c in ("B3", "D3", "F3"): ind[c].font = font()
ind.conditional_formatting.add("D3", CellIsRule(operator="equal", formula=['"PENDIENTE"'], fill=PatternFill("solid", fgColor="F8CBAD"), font=font(bold=True, color="C00000")))
ind.conditional_formatting.add("D3", CellIsRule(operator="equal", formula=['"OK - cargado"'], fill=PatternFill("solid", fgColor="C6EFCE")))
header(ind, 5, ["SKU", "Descripción", "Stock actual", "Consumo último intervalo", "Prom. diario últimos N días",
                "Prom. diario mes en curso", "Consumo mensual promedio", "Días de cobertura", "Estado"])
ind.row_dimensions[5].height = 32
LAST = "$B$3"; NREG = "$F$3"
MONTH0 = 13
months = [dt.date(2026 + (9 + k) // 12, (9 + k) % 12 + 1, 1) for k in range(24)]  # Oct-2026 .. Sep-2028
MR0, MR1 = MONTH0 + 1, MONTH0 + len(months)
for i, (s, rcol, c) in enumerate(blocks):
    r = 6 + i
    ccol = f"Consumo_Diario!${c}${N0}:${c}${N1}"; scol = f"Consumo_Diario!${s}${N0}:${s}${N1}"
    win = f'{FD},">"&({LAST}-{CFG_DIAS}),{FD},"<="&{LAST}'
    mon = f'{FD},">="&DATE(YEAR({LAST}),MONTH({LAST}),1),{FD},"<="&{LAST}'
    ind[f"A{r}"] = f"=Config!$A${SKU_R[i]}"
    ind[f"B{r}"] = f"=Config!$B${SKU_R[i]}"
    ind[f"C{r}"] = f'=IF({NREG}=0,"",INDEX({scol},{NREG}))'
    ind[f"D{r}"] = f'=IF({NREG}<2,"",INDEX({ccol},{NREG}))'
    ind[f"E{r}"] = f'=IF({NREG}<2,"",IFERROR(SUMIFS({ccol},{win})/SUMIFS({DD},{win}),""))'
    ind[f"F{r}"] = f'=IF({NREG}<2,"",IFERROR(SUMIFS({ccol},{mon})/SUMIFS({DD},{mon}),""))'
    mcol = "CFI"[i]
    ind[f"G{r}"] = f'=IFERROR(AVERAGEIFS(${mcol}${MR0}:${mcol}${MR1},$B${MR0}:$B${MR1},">0"),"")'
    ind[f"H{r}"] = f'=IF(OR(C{r}="",E{r}="",E{r}=0),"",C{r}/E{r})'
    ind[f"I{r}"] = f'=IF(H{r}="","Sin datos",IF(H{r}<{CFG_COB},"REPONER","OK"))'
    for col in "ABCDEFGHI": ind[f"{col}{r}"].font = font(); ind[f"{col}{r}"].border = BOX
    for col in "ACDG": ind[f"{col}{r}"].number_format = "#,##0"
    ind[f"C{r}"].number_format = "#,##0"
    for col in "EF": ind[f"{col}{r}"].number_format = "#,##0.0"
    ind[f"H{r}"].number_format = "0.0"
    for col in "AB": ind[f"{col}{r}"].font = GREEN
wb.defined_names["ind_Cobertura"] = DefinedName("ind_Cobertura", attr_text="Indicador!$H$6:$H$8")
ind.conditional_formatting.add("I6:I8", CellIsRule(operator="equal", formula=['"REPONER"'], fill=PatternFill("solid", fgColor="F8CBAD"), font=font(bold=True, color="C00000")))
ind.conditional_formatting.add("I6:I8", CellIsRule(operator="equal", formula=['"OK"'], fill=PatternFill("solid", fgColor="C6EFCE")))

# Tabla mensual
ind["A11"] = "Consumo y promedio mensual"; ind["A11"].font = font(bold=True, size=12, color="1F3864")
ind.merge_cells("A12:B12")
for i in range(3):
    c0 = 3 + 3 * i
    ind.merge_cells(start_row=12, start_column=c0, end_row=12, end_column=c0 + 2)
    cell = ind.cell(row=12, column=c0, value=f"=Config!$A${SKU_R[i]}")
    cell.font = font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="2F5597"); cell.alignment = Alignment(horizontal="center")
header(ind, MONTH0, ["Mes", "Días con datos", "Consumo", "Recepción", "Prom. diario"] + ["Consumo", "Recepción", "Prom. diario"] * 2 + ["Consumo total"])
for k, m in enumerate(months):
    r = MR0 + k
    ind[f"A{r}"] = m; ind[f"A{r}"].number_format = "mmm-yyyy"
    rng = f'{FD},">="&$A{r},{FD},"<"&DATE(YEAR($A{r}),MONTH($A{r})+1,1)'
    ind[f"B{r}"] = f"=SUMIFS({DD},{rng})"
    for i, (s, rcol, c) in enumerate(blocks):
        cc, rr_, pp = [chr(ord("C") + 3 * i + j) for j in range(3)]
        ind[f"{cc}{r}"] = f"=SUMIFS(Consumo_Diario!${c}${N0}:${c}${N1},{rng})"
        ind[f"{rr_}{r}"] = f"=SUMIFS(Consumo_Diario!${rcol}${N0}:${rcol}${N1},{rng})"
        ind[f"{pp}{r}"] = f'=IF($B{r}=0,"",{cc}{r}/$B{r})'
        ind[f"{cc}{r}"].number_format = "#,##0"; ind[f"{rr_}{r}"].number_format = "#,##0"; ind[f"{pp}{r}"].number_format = "#,##0.0"
    ind[f"L{r}"] = f"=C{r}+F{r}+I{r}"; ind[f"L{r}"].number_format = "#,##0"
    for col in "ABCDEFGHIJKL": ind[f"{col}{r}"].font = font(); ind[f"{col}{r}"].border = BOX
avg = MR1 + 1
ind[f"A{avg}"] = "Promedio mensual (meses con datos)"
for col in "CDFGIJL":
    ind[f"{col}{avg}"] = f'=IFERROR(AVERAGEIFS({col}{MR0}:{col}{MR1},$B${MR0}:$B${MR1},">0"),"")'
    ind[f"{col}{avg}"].number_format = "#,##0"
for col in "EHK":
    src = chr(ord(col) - 2)
    ind[f"{col}{avg}"] = f'=IFERROR(SUM({src}{MR0}:{src}{MR1})/SUM($B${MR0}:$B${MR1}),"")'
    ind[f"{col}{avg}"].number_format = "#,##0.0"
for col in "ABCDEFGHIJKL":
    ind[f"{col}{avg}"].font = font(bold=True); ind[f"{col}{avg}"].fill = SUB; ind[f"{col}{avg}"].border = BOX
ind[f"A{avg+1}"] = "Nota: el mes en curso y el primer mes son parciales; use 'Prom. diario' para comparar meses."
ind[f"A{avg+1}"].font = font(italic=True, size=9, color="595959")
ind.column_dimensions["A"].width = 16; ind.column_dimensions["B"].width = 36
for col in "CDEFGHIJKL": ind.column_dimensions[col].width = 14
ind.freeze_panes = "A6"

ch = BarChart(); ch.type = "col"; ch.grouping = "clustered"; ch.title = "Consumo mensual por SKU"
ch.y_axis.title = "Unidades"; ch.height = 8; ch.width = 22
for i in range(3):
    col = 3 + 3 * i
    ch.add_data(Reference(ind, min_col=col, min_row=MONTH0, max_row=MR1), titles_from_data=False)
    ch.series[-1].tx = None
from openpyxl.chart.series import SeriesLabel
for i, s in enumerate(ch.series): s.tx = SeriesLabel(v=str(SKUS[i][0]))
ch.set_categories(Reference(ind, min_col=1, min_row=MR0, max_row=MR1))
# data refs start at header row -> rebuild with min_row=MR0
ch.series.clear()
for i in range(3):
    col = 3 + 3 * i
    ch.add_data(Reference(ind, min_col=col, min_row=MR0, max_row=MR1), titles_from_data=False)
    ch.series[-1].tx = SeriesLabel(v=str(SKUS[i][0]))
ch.set_categories(Reference(ind, min_col=1, min_row=MR0, max_row=MR1))
ind.add_chart(ch, f"A{avg+3}")

# ---------------- Log ----------------
lg = wb.create_sheet("Log")
title(lg, "Bitácora del script", "La escribe el script (más reciente abajo).")
header(lg, 4, ["Fecha y hora", "Evento", "Detalle"])
lg["A5"] = dt.datetime(2026, 10, 15, 8, 45); lg["B5"] = "Carga inicial"; lg["C5"] = "Plantilla creada con el archivo TotalFotoRsp_-_2026101508.xlsx (153523=6725, 153553=5025, 153526=11025)."
lg["A5"].number_format = "dd/mm/yyyy hh:mm"
for col in "ABC": lg[f"{col}5"].font = font()
lg.column_dimensions["A"].width = 18; lg.column_dimensions["B"].width = 22; lg.column_dimensions["C"].width = 110
lg.freeze_panes = "A5"

wb.move_sheet("Indicador", offset=-4)
wb.active = wb.sheetnames.index("Indicador")
wb.save(out)
print("ok", out, "monthly rows", MR0, MR1)
